"""Импорт производственных данных из файлов в формате таблиц кейса.

Поддерживаются DOCX (файл кейса) и XLSX (выгрузка MES/1С или шаблон /api/data/template).
Таблица распознаётся по заголовкам, порядок таблиц и листов не важен.
"""
import io
import re
import zipfile
from datetime import date, datetime
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook

from ..config import CASE_FILE

# заголовки по таблицам: ключ поля → слова, по которым узнаём колонку
SCHEMAS = {
    "lines": {"date": "дата", "line": "линия", "plan": "план", "fact": "факт", "hours": "время работы", "load": "загрузка"},
    "downtime": {"date": "дата", "area": "участок", "equipment": "оборудование", "reason": "причина", "minutes": "длительность"},
    "quality": {"date": "дата", "area": "участок", "output": "выпущено", "defects": "брак"},
    "monthPlan": {"model": "модель", "plan": "план на месяц"},
}
# порядок проверки: у «качества» и «простоев» общие колонки, поэтому сначала более специфичные
ORDER = ("downtime", "quality", "lines", "monthPlan")
SHEET_TITLES = {"lines": "Работа линий", "downtime": "Простои", "quality": "Качество", "monthPlan": "План месяца"}
TEMPLATE_HEADERS = {
    "lines": ["Дата", "Линия", "План", "Факт", "Время работы, ч", "Загрузка, %"],
    "downtime": ["Дата", "Участок", "Оборудование", "Причина", "Длительность, мин"],
    "quality": ["Дата", "Участок", "Выпущено", "Брак", "% брака"],
    "monthPlan": ["Модель", "План на месяц"],
}


class ImportError_(ValueError):
    pass


def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def _date(v) -> str:
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    raise ValueError(f"дата «{s}» не распознана")


def _num(v) -> float:
    if isinstance(v, (int, float)):
        return float(v)
    return float(str(v).replace(" ", "").replace(" ", "").replace(",", "."))


def _columns(header: list, schema: dict) -> dict | None:
    """Индексы колонок по заголовку или None, если таблица не подходит под схему."""
    cells = [_norm(h) for h in header]
    cols = {}
    for key, word in schema.items():
        # точное начало заголовка важнее вхождения: «брак» ≠ «% брака», «план» ≠ «план на месяц»
        idx = next((i for i, c in enumerate(cells) if c == word or c.startswith(word + ",")), None)
        if idx is None:
            idx = next((i for i, c in enumerate(cells) if c.startswith(word) and i not in cols.values()), None)
        if idx is None:
            return None
        cols[key] = idx
    return cols


def parse_tables(tables: list[list[list]]) -> tuple[dict, list[str]]:
    """Таблицы (строки × ячейки) → пакет для хранилища и список ошибок."""
    batch, errors = {}, []
    for table in tables:
        if len(table) < 2:
            continue
        for name in ORDER:
            if name in batch:
                continue
            cols = _columns(table[0], SCHEMAS[name])
            if cols is None:
                continue
            rows = []
            for n, raw in enumerate(table[1:], start=2):
                if all(_norm(c) == "" for c in raw):
                    continue
                try:
                    rows.append(_row(name, {k: raw[i] if i < len(raw) else "" for k, i in cols.items()}))
                except (ValueError, TypeError) as e:
                    errors.append(f"{SHEET_TITLES[name]}, строка {n}: {e}")
            batch[name] = rows
            break
    if not batch:
        errors.append("Не найдено ни одной таблицы с заголовками кейса. Скачайте шаблон и заполните его.")
    errors += _check(batch)
    return batch, errors


def _row(name: str, r: dict) -> dict:
    if name == "lines":
        line = str(r["line"]).strip()
        return {"date": _date(r["date"]), "line": line, "area": re.split(r"[-\s]", line)[0],
                "plan": int(_num(r["plan"])), "fact": int(_num(r["fact"])), "hours": _num(r["hours"]),
                "load": _num(r["load"])}
    if name == "downtime":
        return {"date": _date(r["date"]), "area": str(r["area"]).strip(), "equipment": str(r["equipment"]).strip(),
                "reason": str(r["reason"]).strip(), "minutes": int(_num(r["minutes"]))}
    if name == "quality":
        return {"date": _date(r["date"]), "area": str(r["area"]).strip(), "output": int(_num(r["output"])),
                "defects": int(_num(r["defects"]))}
    return {"model": str(r["model"]).strip(), "plan": int(_num(r["plan"]))}


def _check(batch: dict) -> list[str]:
    """Расчёт OEE сопоставляет строку линии с качеством того же участка за ту же дату."""
    quality = {(q["date"], q["area"]) for q in batch.get("quality", [])}
    if "lines" in batch and "quality" in batch:
        return [f"Нет качества для {r['line']} за {r['date']}" for r in batch["lines"] if (r["date"], r["area"]) not in quality]
    if "lines" in batch:
        return ["В файле есть работа линий, но нет таблицы качества — загрузите их вместе."]
    return []


def docx_tables(data: bytes) -> list[list[list]]:
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    root = ElementTree.fromstring(zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml"))
    return [[["".join(t.text or "" for t in tc.iter(f"{w}t")) for tc in tr.iter(f"{w}tc")]
             for tr in tbl.iter(f"{w}tr")] for tbl in root.iter(f"{w}tbl")]


def xlsx_tables(data: bytes) -> list[list[list]]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    tables = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        # заголовок — первая строка, где есть хотя бы две заполненные ячейки
        start = next((i for i, r in enumerate(rows) if sum(c not in (None, "") for c in r) >= 2), None)
        if start is not None:
            tables.append([["" if c is None else c for c in r] for r in rows[start:]])
    return tables


def read(filename: str, data: bytes) -> tuple[dict, list[str]]:
    ext = Path(filename).suffix.lower()
    try:
        if ext == ".docx":
            return parse_tables(docx_tables(data))
        if ext in (".xlsx", ".xlsm"):
            return parse_tables(xlsx_tables(data))
    except (zipfile.BadZipFile, KeyError, ElementTree.ParseError) as e:
        return {}, [f"Файл повреждён или это не {ext}: {e}"]
    return {}, [f"Формат {ext or 'без расширения'} не поддерживается: загрузите .xlsx или .docx"]


def case_dataset() -> dict:
    batch, errors = read(CASE_FILE.name, CASE_FILE.read_bytes())
    if errors:
        raise ImportError_("; ".join(errors))
    return batch


def template(ds: dict) -> bytes:
    """XLSX с четырьмя листами в формате кейса, заполненный текущими данными."""
    wb = Workbook()
    wb.remove(wb.active)
    for name, headers in TEMPLATE_HEADERS.items():
        ws = wb.create_sheet(SHEET_TITLES[name])
        ws.append(headers)
        for r in ds.get(name, []):
            d = datetime.strptime(r["date"], "%Y-%m-%d").strftime("%d.%m.%Y") if "date" in r else None
            ws.append({
                "lines": lambda: [d, r["line"], r["plan"], r["fact"], r["hours"], r["load"]],
                "downtime": lambda: [d, r["area"], r["equipment"], r["reason"], r["minutes"]],
                "quality": lambda: [d, r["area"], r["output"], r["defects"], round(r["defects"] / r["output"] * 100, 1) if r["output"] else 0],
                "monthPlan": lambda: [r["model"], r["plan"]],
            }[name]())
        for col, h in zip("ABCDEF", headers):
            ws.column_dimensions[col].width = max(14, len(h) + 4)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
